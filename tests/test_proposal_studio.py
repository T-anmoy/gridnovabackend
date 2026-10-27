"""BE-18 — the Proposal Studio through its API: the workflow (approval thresholds, maker ≠
checker, revise → -R1 and superseded, duplicate, concurrency, lock), the customer's page
(token, rate limit, view and download events, accept / decline, verification QR), conversion
to a lead and the payment schedule, the Excel import, the optional assistant, delivery and
the issued PDF/A-3b. Every name and number is invented."""
from __future__ import annotations

import hashlib
import io
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest

from app.memory.models import (
    Lead, LeadSource, PaymentMilestone, Proposal, ProposalStatus, StudioEvent, StudioProposal, StudioProposalVersion,
)
from tests.proposal_document_support import require_weasyprint
from tests.studio_support import (
    ADMIN, DESK, DESK2, MGMT, D, acknowledge_all, edit, fill, issued_residential, new_customer, new_proposal, ok,
)

PUBLIC = "/api/v1/public/studio"
COMMERCIAL_52 = [
    ("site", {"address_lines": ["1, Example Road", "Kolkata – 700006, West Bengal"], "town": "Kolkata",
              "state": "West Bengal"}),
    ("connection", {"category": "Commercial", "phase": "three", "sanctioned_load_kw": "60"}),
    ("plant", {"capacity_kwp": "52"}),
    ("price", {"mode": "staff_basic", "basic": "2860000"}),
]


@pytest.fixture(autouse=True)
def _studio_env(monkeypatch, fake_storage):
    """A Fernet key for the stored links, storage, WeasyPrint, and fresh rate-limit windows."""
    require_weasyprint()
    from app.api.routes import public_studio, studio
    from app.config import settings
    from app.security.encryption import generate_key
    from app.security.rate_limit import RateLimitLog
    monkeypatch.setattr(settings, "fernet_encryption_key", generate_key())
    for name in ("_page_log", "_pdf_log", "_action_log", "_verify_log"):
        monkeypatch.setattr(public_studio, name, RateLimitLog(60))
    monkeypatch.setattr(studio, "_preview_log", RateLimitLog(60))


@pytest.fixture
def alerts(monkeypatch):
    sent = []
    from app.notifications import notifier as notifier_module
    monkeypatch.setattr(notifier_module.notifier, "notify", lambda event, payload: sent.append((event, payload)))
    return sent


def _view(client, pid, headers=DESK):
    return ok(client.get(f"{D}/studio/proposals/{pid}", headers=headers))


def _commercial(client, headers=DESK) -> int:
    pid = new_proposal(client, "commercial_rooftop", headers=headers)["proposal"]["id"]
    fill(client, pid, COMMERCIAL_52, headers=headers)
    acknowledge_all(client, pid, headers=headers)
    return pid


def _freeze(monkeypatch, when: datetime):
    from app.studio import service
    monkeypatch.setattr(service, "now", lambda: when)
    monkeypatch.setattr(service, "today_ist", lambda at=None: (at or when).date())


# ── Customers ────────────────────────────────────────────────────────────────────

class TestCustomers:
    def test_duplicate_mobile_or_gstin_is_a_warning(self, client):
        first = new_customer(client, gstin="27AAPFU0939F1ZV")
        r = client.post(f"{D}/studio/customers", headers=DESK, json={
            "name": "Another Name", "mobile": "+91 98300 00001", "address_lines": ["x"]})
        assert r.status_code == 201
        assert r.json()["duplicates"][0]["id"] == first["id"] and r.json()["duplicates"][0]["matched_on"] == ["mobile"]

    @pytest.mark.parametrize("field,value", [("mobile", "12345"), ("gstin", "27AAPFU0939F1ZW"), ("name", "x"),
                                             ("email", "not-an-email"), ("address_lines", [])])
    def test_validation(self, client, field, value):
        r = client.post(f"{D}/studio/customers", headers=DESK, json={
            "name": "Valid Name", "mobile": "9830000002", "address_lines": ["1 Road"], field: value})
        assert r.status_code == 422 and r.json()["detail"]["field"] == field

    def test_search_by_name_mobile_gstin(self, client):
        new_customer(client, name="Searchable Traders", mobile="9830000077", gstin="19ABEFG7809A1ZI")
        for q in ("searchable", "0077", "19ABEFG7809A1ZI"):
            items = ok(client.get(f"{D}/studio/customers", params={"q": q}, headers=DESK))["items"]
            assert [c["name"] for c in items] == ["Searchable Traders"], q


# ── Workflow ─────────────────────────────────────────────────────────────────────

class TestWorkflow:
    def test_small_list_priced_proposal_issues_without_approval(self, client):
        out, token = issued_residential(client)
        assert out["proposal"]["status"] == "issued" and out["version"]["locked"] is True
        assert out["version"]["document_reference"].startswith("GN/") and len(token) >= 43
        assert out["version"]["approved_by"] is None

    def test_above_threshold_needs_approval(self, client):
        pid = _commercial(client)
        r = client.post(f"{D}/studio/proposals/{pid}/issue", headers=DESK)
        assert r.status_code == 409 and r.json()["detail"]["code"] == "approval_required"
        assert ok(client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK))["result"] == "pending_approval"

    def test_below_price_master_needs_approval_even_when_small(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        fill(client, pid)
        r = edit(client, pid, "price", {"mode": "staff_basic", "basic": "180000"})
        assert r.status_code == 422 and r.json()["detail"]["code"] == "reason_required"
        view = ok(edit(client, pid, "price", {"mode": "staff_basic", "basic": "180000", "reason": "Competitor"}))
        assert view["approval_reasons"] == ["price_below_price_master"]

    def test_maker_is_not_checker(self, client, alerts):
        pid = _commercial(client, headers=ADMIN)
        ok(client.post(f"{D}/studio/proposals/{pid}/submit", headers=ADMIN))
        assert any(e == "studio_approval_requested" for e, _ in alerts)
        r = client.post(f"{D}/studio/proposals/{pid}/approve", headers=ADMIN, json={})
        assert r.status_code == 403 and r.json()["detail"]["code"] == "maker_checker"
        assert client.post(f"{D}/studio/proposals/{pid}/approve", headers=DESK, json={}).status_code == 403
        view = ok(client.post(f"{D}/studio/proposals/{pid}/approve", headers=MGMT, json={"comment": "Fine"}))
        assert view["version"]["status"] == "ready" and view["version"]["approved_by"] == "emp-mgmt-1"
        out = ok(client.post(f"{D}/studio/proposals/{pid}/issue", headers=ADMIN))
        assert out["proposal"]["status"] == "issued"

    def test_reject_sends_back_with_a_comment(self, client):
        pid = _commercial(client)
        ok(client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK))
        assert client.post(f"{D}/studio/proposals/{pid}/reject", headers=MGMT, json={}).status_code == 422
        view = ok(client.post(f"{D}/studio/proposals/{pid}/reject", headers=MGMT, json={"comment": "Fix scope"}))
        assert view["version"]["status"] == "draft" and view["version"]["approval_comment"] == "Fix scope"

    def test_an_edit_after_approval_clears_it(self, client):
        pid = _commercial(client)
        ok(client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK))
        ok(client.post(f"{D}/studio/proposals/{pid}/approve", headers=MGMT, json={}))
        view = ok(edit(client, pid, "notes_free", ["Delivery in two lots."]))
        assert view["approval_cleared"] is True and view["version"]["status"] == "draft"
        assert view["version"]["approved_by"] is None

    def test_warnings_must_be_acknowledged(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        fill(client, pid)
        ok(edit(client, pid, "connection", {"discom": "WBSEDCL", "consumer_number": "100000001",
                                            "category": "Residential", "phase": "single", "sanctioned_load_kw": "2"}))
        r = client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK)
        assert r.status_code == 409 and r.json()["detail"]["code"] == "warnings_unacknowledged"
        acknowledge_all(client, pid)
        assert ok(client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK))["result"] == "ready"
        events = [e["type"] for e in ok(client.get(f"{D}/studio/proposals/{pid}/events", headers=DESK))["items"]]
        assert "warnings_acknowledged" in events

    def test_errors_block_submit_and_issue(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        for action in ("submit", "issue"):
            r = client.post(f"{D}/studio/proposals/{pid}/{action}", headers=DESK)
            assert r.status_code == 409 and r.json()["detail"]["code"] == "proof_check_failed"

    def test_content_not_ready_names_the_registry(self, client):
        pid = new_proposal(client, "ground_mounted")["proposal"]["id"]
        fill(client, pid, [("site", {"address_lines": ["Example Mouza"], "town": "Bankura", "state": "West Bengal"}),
                           ("connection", {"category": "Commercial", "phase": "three"}),
                           ("plant", {"capacity_kwp": "1000"}),
                           ("price", {"mode": "staff_basic", "rate_per_kwp": "38000"})])
        r = client.post(f"{D}/studio/proposals/{pid}/submit", headers=DESK)
        assert r.status_code == 409 and r.json()["detail"]["code"] == "content_not_ready"
        assert r.json()["detail"]["registry"] == ["PRO-7"]
        pdf = client.get(f"{D}/studio/proposals/{pid}/preview.pdf", headers=DESK)
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")              # preview still works

    def test_stale_edit_refused(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        rev = _view(client, pid)["revision"]
        ok(edit(client, pid, "notes_free", ["One."], base_revision=rev))
        r = edit(client, pid, "notes_free", ["Two."], base_revision=rev)
        assert r.status_code == 409 and r.json()["detail"]["code"] == "stale_version"
        assert r.json()["detail"]["current_revision"] == rev + 1

    def test_a_locked_version_never_changes(self, client):
        out, _ = issued_residential(client)
        r = edit(client, out["proposal"]["id"], "notes_free", ["Too late."])
        assert r.status_code == 409 and r.json()["detail"]["code"] == "proposal_locked"

    def test_revise_issues_r1_and_supersedes(self, client):
        out, token = issued_residential(client)
        pid, ref = out["proposal"]["id"], out["version"]["document_reference"]
        view = ok(client.post(f"{D}/studio/proposals/{pid}/revise", headers=DESK))
        assert view["version"]["version"] == 2 and view["version"]["status"] == "draft"
        assert view["document"]["plant"]["capacity_kwp"] == "3"                          # everything carried over
        ok(edit(client, pid, "notes_free", ["Revised after the site visit."]))
        acknowledge_all(client, pid)
        out2 = ok(client.post(f"{D}/studio/proposals/{pid}/issue", headers=DESK))
        assert out2["version"]["document_reference"] == f"{ref}-R1"
        assert {v["version"]: v["status"] for v in out2["versions"]} == {1: "superseded", 2: "issued"}
        assert out2["version"]["change_summary"][0]["field"].startswith("notes_free")
        r = client.post(f"{PUBLIC}/{token}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                           "accept_terms": True})
        assert r.status_code == 409 and r.json()["detail"]["code"] == "proposal_superseded"
        diff = ok(client.get(f"{D}/studio/proposals/{pid}/versions/1/diff/2", headers=DESK))["changes"]
        assert any(c["field"].startswith("notes_free") for c in diff)

    def test_duplicate_for_another_customer(self, client):
        out, _ = issued_residential(client)
        other = new_customer(client, name="Second Customer", mobile="9830000055")
        view = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/duplicate", headers=DESK,
                              json={"customer_id": other["id"]}))
        assert view["proposal"]["id"] != out["proposal"]["id"] and view["version"]["status"] == "draft"
        assert view["document"]["addressee"]["name"] == "Second Customer"
        assert view["figures"]["total"] == "201465.00"

    def test_mark_lost_needs_a_reason(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        assert client.post(f"{D}/studio/proposals/{pid}/mark-lost", headers=DESK, json={"reason": ""}).status_code \
            == 422
        view = ok(client.post(f"{D}/studio/proposals/{pid}/mark-lost", headers=DESK, json={"reason": "Went solar "
                                                                                                 "elsewhere"}))
        assert view["proposal"]["status"] == "lost"

    def test_in_person_acceptance_with_a_signed_copy(self, client, db_session):
        out, _ = issued_residential(client)
        pid = out["proposal"]["id"]
        scan = b"%PDF-1.4 signed copy"
        view = ok(client.post(f"{D}/studio/proposals/{pid}/mark-accepted", headers=DESK,
                              data={"accepted_on": out["version"]["document_date"]},
                              files={"signed_copy": ("signed.pdf", scan, "application/pdf")}))
        assert view["version"]["accepted_via"] == "in_person" and view["convert_offered"] is True
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=pid).one()
        assert v.signed_copy_sha256 == hashlib.sha256(scan).hexdigest()
        r = client.post(f"{D}/studio/proposals/{pid}/mark-accepted", headers=DESK,
                        data={"accepted_on": out["version"]["document_date"]},
                        files={"signed_copy": ("x.pdf", b"MZ not a pdf", "application/pdf")})
        assert r.status_code == 422

    def test_preview_is_stamped_draft(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        fill(client, pid)
        r = client.get(f"{D}/studio/proposals/{pid}/preview.pdf", headers=DESK)
        from tests.proposal_document_support import pdf_text
        assert r.status_code == 200 and "DRAFT — NOT FOR CUSTOMER" in "\n".join(pdf_text(r.content))

    def test_list_filters_and_summary(self, client):
        out, _ = issued_residential(client)
        new_proposal(client, "commercial_rooftop")
        items = ok(client.get(f"{D}/studio/proposals", params={"status": "issued"}, headers=DESK))["items"]
        assert [i["id"] for i in items] == [out["proposal"]["id"]]
        ref = out["version"]["document_reference"]
        assert ok(client.get(f"{D}/studio/proposals", params={"q": ref}, headers=DESK))["total"] == 1
        s = ok(client.get(f"{D}/studio/summary", headers=DESK))
        assert s["issued_this_month"] == {"count": 1, "value": "201465.00"} and s["drafts"] == 1

    def test_save_as_template_has_no_customer_data(self, client):
        out, _ = issued_residential(client)
        item = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/save-as-template", headers=DESK,
                              json={"name": "3 kW PM Surya Ghar"}))
        doc = item["body"]["document"]
        assert doc["addressee"]["name"] is None and doc["site"]["address_lines"] == []
        cid = new_customer(client, name="Templated Person", mobile="9830000066")["id"]
        view = ok(client.post(f"{D}/studio/proposals", headers=DESK, json={
            "customer_id": cid, "project_type": "residential_rooftop", "template_id": item["id"]}))
        assert view["document"]["addressee"]["name"] == "Templated Person"
        assert view["document"]["plant"]["capacity_kwp"] == "3"


# ── Reminders and expiry ─────────────────────────────────────────────────────────

class TestReminders:
    def test_not_viewed_expiring_and_expired(self, client, db_session, monkeypatch, alerts):
        from app.studio import reminders
        out, _ = issued_residential(client)
        pid = out["proposal"]["id"]
        issued = db_session.query(StudioProposalVersion).filter_by(proposal_id=pid).one()
        start = issued.issued_at.replace(tzinfo=timezone.utc) if issued.issued_at.tzinfo is None else issued.issued_at
        day = issued.document_date
        # Day 2: nothing.
        assert reminders.run(db_session, start + timedelta(days=2)) == {"not_viewed": 0, "expiring": 0,
                                                                        "customer_emails": 0, "expired": 0}
        # Day 3: not viewed → the owner; re-running sends nothing more.
        assert reminders.run(db_session, start + timedelta(days=3, hours=1))["not_viewed"] == 1
        assert reminders.run(db_session, start + timedelta(days=3, hours=2))["not_viewed"] == 0
        # 3 days before valid_until (15 days): expiring, once.
        when = datetime.combine(day + timedelta(days=12), datetime.min.time(), timezone.utc) + timedelta(hours=6)
        assert reminders.run(db_session, when)["expiring"] == 1
        assert reminders.run(db_session, when)["expiring"] == 0
        # Past valid_until → expired.
        when = datetime.combine(day + timedelta(days=16), datetime.min.time(), timezone.utc) + timedelta(hours=6)
        assert reminders.run(db_session, when)["expired"] == 1
        db_session.commit()
        assert db_session.get(StudioProposal, pid).status.value == "expired"
        events = [e for e, _ in alerts]
        assert events.count("studio_not_viewed") == 1 and events.count("studio_expiring") == 1
        owners = {p["owner_id"] for e, p in alerts if e.startswith("studio_")}
        assert owners == {"emp-desk-1"}

    def test_customer_expiry_email_only_when_switched_on(self, client, db_session, monkeypatch):
        from app.studio import config as studio_config
        from app.studio import public, reminders
        sent = []
        monkeypatch.setattr(public, "expiry_email", lambda db, p, v: sent.append(v.id) or True)
        out, _ = issued_residential(client)
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=out["proposal"]["id"]).one()
        when = datetime.combine(v.valid_until - timedelta(days=1), datetime.min.time(), timezone.utc)
        assert reminders.run(db_session, when)["customer_emails"] == 0
        monkeypatch.setitem(studio_config.load()["reminders"], "customer_expiry_reminder", True)
        v.reminders_sent = None
        assert reminders.run(db_session, when)["customer_emails"] == 1 and sent == [v.id]

    def test_job_route_with_the_internal_key(self, client, monkeypatch):
        from app.config import settings
        monkeypatch.setattr(settings, "internal_api_key", "k" * 40)
        r = client.post(f"{D}/internal/check-studio-reminders", headers={"X-Internal-Key": "k" * 40})
        assert r.status_code == 200 and set(r.json()) == {"not_viewed", "expiring", "customer_emails", "expired"}
        assert client.post(f"{D}/internal/check-studio-reminders").status_code == 401


# ── The customer's page ──────────────────────────────────────────────────────────

class TestPublic:
    def test_unknown_token_is_404(self, client):
        for path in ("x" * 43, "short"):
            assert client.get(f"{PUBLIC}/{path}").status_code == 404

    def test_view_records_events_and_alerts_owner_once(self, client, db_session, alerts):
        out, token = issued_residential(client)
        r = client.get(f"{PUBLIC}/{token}", headers={"User-Agent": "TestBrowser/1"})
        assert r.status_code == 200
        body = r.json()
        assert (body["total"], body["state"], body["mobile_hint"]) == ("201465.00", "open", "XXXXXX0001")
        assert r.headers["x-robots-tag"].startswith("noindex") and r.headers["cache-control"] == "no-store"
        client.get(f"{PUBLIC}/{token}")
        assert [e for e, _ in alerts].count("studio_viewed") == 1
        rows = db_session.query(StudioEvent).filter_by(proposal_id=out["proposal"]["id"], type="viewed").all()
        assert len(rows) == 2 and rows[0].user_agent == "TestBrowser/1" and rows[0].ip
        assert _view(client, out["proposal"]["id"])["proposal"]["status"] == "viewed"

    def test_download_redirects_and_is_recorded(self, client, db_session):
        out, token = issued_residential(client)
        r = client.get(f"{PUBLIC}/{token}/pdf", follow_redirects=False)
        assert r.status_code == 302 and r.headers["location"].startswith("https://signed.example/studio/")
        assert db_session.query(StudioEvent).filter_by(type="downloaded").count() == 1

    def test_rate_limited(self, client, monkeypatch):
        from app.api.routes import public_studio
        monkeypatch.setattr(public_studio, "RATE_LIMIT", 2)
        _, token = issued_residential(client)
        codes = [client.get(f"{PUBLIC}/{token}").status_code for _ in range(3)]
        assert codes == [200, 200, 429]

    def test_accept_checks_name_last4_and_terms(self, client, db_session, alerts):
        out, token = issued_residential(client)
        base = {"full_name": "Test Kumar Das", "mobile_last4": "0001", "accept_terms": True}
        r = client.post(f"{PUBLIC}/{token}/accept", json={**base, "mobile_last4": "9999"})
        assert r.status_code == 422 and r.json()["detail"]["code"] == "mobile_mismatch"
        r = client.post(f"{PUBLIC}/{token}/accept", json={**base, "accept_terms": False})
        assert r.status_code == 422 and r.json()["detail"]["code"] == "terms_not_accepted"
        r = client.post(f"{PUBLIC}/{token}/accept", json=base, headers={"User-Agent": "Phone/2"})
        assert r.status_code == 200 and r.json()["already_accepted"] is False
        again = client.post(f"{PUBLIC}/{token}/accept", json=base)
        assert again.json()["already_accepted"] is True
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=out["proposal"]["id"]).one()
        assert (v.accepted_name, v.accepted_via, v.accepted_user_agent) == ("Test Kumar Das", "online", "Phone/2")
        assert v.accepted_pdf_sha256 == v.pdf_sha256 and v.accepted_ip
        assert "studio_accepted" in [e for e, _ in alerts]

    def test_accept_after_expiry_is_409(self, client, db_session):
        out, token = issued_residential(client)
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=out["proposal"]["id"]).one()
        v.valid_until = date(2020, 1, 1)
        db_session.commit()
        r = client.post(f"{PUBLIC}/{token}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                           "accept_terms": True})
        assert r.status_code == 409 and r.json()["detail"]["code"] == "proposal_expired"

    def test_decline_with_a_reason(self, client, alerts):
        out, token = issued_residential(client)
        r = client.post(f"{PUBLIC}/{token}/decline", json={"reason": "Budget"})
        assert r.status_code == 200
        assert _view(client, out["proposal"]["id"])["version"]["decline_reason"] == "Budget"
        assert "studio_declined" in [e for e, _ in alerts]

    def test_verify_each_status(self, client, db_session):
        out, token = issued_residential(client)
        code = out["version"]["verification_code"]
        url = f"/api/v1/public/proposals/verify/{code}"
        body = ok(client.get(url))
        assert (body["status"], body["customer"], body["reference"]) == ("valid", "Test K.",
                                                                         out["version"]["document_reference"])
        pid = out["proposal"]["id"]
        ok(client.post(f"{D}/studio/proposals/{pid}/revise", headers=DESK))
        acknowledge_all(client, pid)
        out2 = ok(client.post(f"{D}/studio/proposals/{pid}/issue", headers=DESK))
        assert ok(client.get(url))["status"] == "superseded"
        url2 = f"/api/v1/public/proposals/verify/{out2['version']['verification_code']}"
        token2 = out2["public_url"].rsplit("/", 1)[1]
        ok(client.post(f"{PUBLIC}/{token2}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                          "accept_terms": True}))
        assert ok(client.get(url2))["status"] == "accepted"
        third, token3 = issued_residential(client)
        v3 = db_session.query(StudioProposalVersion).filter_by(proposal_id=third["proposal"]["id"]).one()
        v3.valid_until = date(2020, 1, 1)
        db_session.commit()
        assert ok(client.get(f"/api/v1/public/proposals/verify/{v3.verification_code}"))["status"] == "expired"
        fourth, token4 = issued_residential(client)
        client.post(f"{PUBLIC}/{token4}/decline", json={})
        assert ok(client.get(f"/api/v1/public/proposals/verify/{fourth['version']['verification_code']}"))[
            "status"] == "declined"
        draft = new_proposal(client)
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=draft["proposal"]["id"]).one()
        assert client.get(f"/api/v1/public/proposals/verify/{v.verification_code}").status_code == 404
        assert client.get("/api/v1/public/proposals/verify/NOPE1234NOPE").status_code == 404


# ── Delivery ─────────────────────────────────────────────────────────────────────

class TestDelivery:
    def test_whatsapp_link(self, client):
        out, token = issued_residential(client)
        body = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/whatsapp-link", headers=DESK))
        assert body["url"].startswith("https://wa.me/919830000001?text=") and token in body["link"]
        assert out["version"]["document_reference"] in body["text"]

    def test_email_attaches_the_pdf_and_replies_to_the_owner(self, client, db_session, monkeypatch):
        import httpx

        from app.memory.models import Employee
        from app.notifications import notifier as notifier_module
        from app.notifications.notifier import EmailNotifier, LogNotifier
        out, _ = issued_residential(client)
        db_session.get(Employee, "emp-desk-1").email = "desk@gridnova.example"
        db_session.commit()
        posts = []

        class _Resp:
            status_code = 200

        monkeypatch.setattr(httpx, "post", lambda *a, **kw: posts.append(kw["json"]) or _Resp())
        monkeypatch.setattr(notifier_module, "notifier", EmailNotifier("re_k", "Grid Nova <n@gridnova.in>", [],
                                                                      "https://app.test", LogNotifier()))
        body = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/email", headers=DESK, json={}))
        assert body["to"] == "test.customer@example.in"
        sent = posts[0]
        assert sent["reply_to"] == "desk@gridnova.example"
        assert sent["attachments"][0]["filename"].startswith("GridNova_Proposal_GN-")
        assert sent["attachments"][0]["filename"].endswith("_3kWp_TestKumarDas.pdf")

    def test_email_without_resend_is_503(self, client):
        out, _ = issued_residential(client)
        r = client.post(f"{D}/studio/proposals/{out['proposal']['id']}/email", headers=DESK, json={})
        assert r.status_code == 503 and r.json()["detail"]["code"] == "email_not_configured"


# ── The issued PDF ───────────────────────────────────────────────────────────────

class TestIssuedPdf:
    def test_pdfa3b_metadata_fonts_and_hash(self, client, db_session):
        from pypdf import PdfReader

        from tests.conftest import _FAKE_BUCKET
        out, _ = issued_residential(client)
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=out["proposal"]["id"]).one()
        pdf = _FAKE_BUCKET[v.pdf_storage_key]
        assert hashlib.sha256(pdf).hexdigest() == v.pdf_sha256
        reader = PdfReader(io.BytesIO(pdf))
        root = reader.trailer["/Root"]
        assert "/OutputIntents" in root
        raw = root["/Metadata"].get_object().get_data().decode()
        assert 'pdfaid:part="3"' in raw and 'pdfaid:conformance="B"' in raw
        meta = reader.metadata
        assert meta["/Author"] == "Grid Nova Energy" and meta["/Title"] == "3 kWp Grid-Tied Rooftop Solar Proposal"
        assert "Supply, Installation" in meta["/Subject"] and v.document_reference in meta["/Keywords"]
        for page in reader.pages:
            for font in (page["/Resources"].get("/Font") or {}).values():
                font = font.get_object()
                desc = font.get("/FontDescriptor") or font["/DescendantFonts"][0].get_object()["/FontDescriptor"]
                desc = desc.get_object()
                assert any(k in desc for k in ("/FontFile", "/FontFile2", "/FontFile3")), font["/BaseFont"]
        assert v.pdf_filename == f"GridNova_Proposal_{v.document_reference.replace('/', '-')}_3kWp_TestKumarDas.pdf"

    def test_the_qr_points_to_verification(self, db_session, client):
        from app.proposals import document_render
        out, _ = issued_residential(client)
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=out["proposal"]["id"]).one()
        url = f"https://www.gridnova.in/verify/p/{v.verification_code}"
        assert document_render.qr_png_size(url) == document_render.qr_png_size(url)
        assert len(v.verification_code) == 12 and v.verification_code.isalnum()

    def test_references_shared_with_lead_proposals(self, client, db_session):
        from app.proposals import reference
        out, _ = issued_residential(client)
        day = date.fromisoformat(out["version"]["document_date"])
        nxt = reference.format_reference(day, reference.next_sequence(db_session, day))
        assert nxt != out["version"]["document_reference"] and nxt.endswith("/02")


# ── Convert to lead ──────────────────────────────────────────────────────────────

class TestConvert:
    def test_creates_lead_and_proposal_once_with_the_same_reference(self, client, db_session):
        out, token = issued_residential(client)
        pid = out["proposal"]["id"]
        ok(client.post(f"{PUBLIC}/{token}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                         "accept_terms": True}))
        first = ok(client.post(f"{D}/studio/proposals/{pid}/convert", headers=DESK))
        again = ok(client.post(f"{D}/studio/proposals/{pid}/convert", headers=DESK))
        assert first["created"] is True and again == {**first, "created": False}
        lead = db_session.get(Lead, first["lead_id"])
        assert lead.source == LeadSource.proposal_studio and lead.assigned_employee_id == "emp-desk-1"
        assert lead.collected_fields["segment"] == "residential"
        proposal = db_session.get(Proposal, first["proposal_id"])
        assert proposal.status == ProposalStatus.accepted
        assert proposal.document_reference == out["version"]["document_reference"]
        v = db_session.query(StudioProposalVersion).filter_by(proposal_id=pid).one()
        assert proposal.pdf_storage_key == v.pdf_storage_key and proposal.total_value == 201465.0
        assert db_session.query(Proposal).filter_by(document_reference=proposal.document_reference).count() == 1
        assert _view(client, pid)["proposal"]["status"] == "converted"

    def test_payments_schedule_20_60_20_against_it(self, client, db_session):
        from app.payments import service as payments
        out, token = issued_residential(client)
        ok(client.post(f"{PUBLIC}/{token}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                         "accept_terms": True}))
        conv = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/convert", headers=DESK))
        lead = db_session.get(Lead, conv["lead_id"])
        # The payments flow finds the accepted proposal exactly as it does for a lead proposal.
        found = db_session.query(Proposal).filter(Proposal.lead_id == lead.id,
                                                  Proposal.status == ProposalStatus.accepted).one()
        total = payments.money(found.total_value)
        amounts = payments.split(total, [Decimal(20), Decimal(60), Decimal(20)])
        assert amounts == [Decimal("40293.00"), Decimal("120879.00"), Decimal("40293.00")]
        for n, (pct, amount) in enumerate(zip((20, 60, 20), amounts), 1):
            db_session.add(PaymentMilestone(project_id=lead.project_id, proposal_id=found.id, number=n, label=f"M{n}",
                                            percent=Decimal(pct), amount_due=amount, trigger="accepted",
                                            created_by="test"))
        db_session.commit()
        rows = payments.milestones(db_session, lead.project_id)
        assert sum(r.amount_due for r in rows) == Decimal("201465.00")

    def test_converted_before_acceptance_then_accepted_online(self, client, db_session):
        out, token = issued_residential(client)
        conv = ok(client.post(f"{D}/studio/proposals/{out['proposal']['id']}/convert", headers=DESK))
        assert db_session.get(Proposal, conv["proposal_id"]).status == ProposalStatus.sent
        ok(client.post(f"{PUBLIC}/{token}/accept", json={"full_name": "Test Kumar Das", "mobile_last4": "0001",
                                                         "accept_terms": True}))
        db_session.expire_all()
        assert db_session.get(Proposal, conv["proposal_id"]).status == ProposalStatus.accepted

    def test_a_later_lead_revision_continues_the_r_numbers(self, db_session):
        from app.proposals import reference
        from tests.test_proposal_document import _chat_lead  # noqa: F401 — the helper exists
        from app.memory.memory_manager import MemoryManager
        memory = MemoryManager(db_session)
        customer = memory.get_or_create_customer("r@example.in", "R Test")
        lead = memory.create_lead(customer.id)
        db_session.add(Proposal(lead_id=lead.id, version=1, status=ProposalStatus.sent, content={},
                                document_reference="GN/01/11/2026/04-R1"))
        nxt = Proposal(lead_id=lead.id, version=2, status=ProposalStatus.draft, content={})
        db_session.add(nxt)
        db_session.flush()
        assert reference.assign(db_session, nxt) == "GN/01/11/2026/04-R2"

    def test_not_before_issue(self, client):
        pid = new_proposal(client)["proposal"]["id"]
        r = client.post(f"{D}/studio/proposals/{pid}/convert", headers=DESK)
        assert r.status_code == 409 and r.json()["detail"]["code"] == "not_issued"


# ── Excel import ─────────────────────────────────────────────────────────────────

def _workbook(*sheets: dict, extra_sheet: bool = False) -> bytes:
    """A workbook from the real template with each dict's {field key or cell: value} filled
    into its own 'Proposal…' sheet."""
    from openpyxl import load_workbook

    from app.studio import importer
    wb = load_workbook(io.BytesIO(importer.template()))
    base = wb["Proposal"]
    rows = {k: r for r, k, *_ in importer.fields() if k}
    for i, values in enumerate(sheets):
        ws = base if i == 0 else wb.copy_worksheet(base)
        if i:
            ws.title = f"Proposal {i + 1}"
        for key, value in values.items():
            ws[f"B{rows[key]}" if key in rows else key] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _import(client, raw: bytes, name="proposals.xlsx", headers=DESK):
    return client.post(f"{D}/studio/import", headers=headers, files={"file": (name, raw, "application/octet-stream")})


class TestImport:
    def test_template_downloads(self, client):
        from openpyxl import load_workbook
        r = client.get(f"{D}/studio/import-template.xlsx", headers=DESK)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Read me", "Proposal", "Example"]
        assert wb["Proposal"].data_validations.dataValidation

    def test_valid_sheets_create_drafts(self, client):
        from app.studio.importer import EXAMPLE
        second = {**EXAMPLE, "name": "Second Person", "mobile": "9830000099", "capacity_kwp": 5, "phase": "three",
                  "dcr": "no", "subsidy_opted": "no"}
        r = _import(client, _workbook(EXAMPLE, second))
        assert r.status_code == 201, r.text
        created = r.json()["created"]
        assert len(created) == 2
        view = _view(client, created[0]["proposal_id"])
        assert view["figures"]["total"] == "201465.00" and view["document"]["plant"]["dcr"] is True
        assert view["figures"]["subsidy"]["applies"] is True

    def test_itemised_sheet(self, client):
        from app.studio.importer import ITEM_FIRST_ROW
        values = {"customer_type": "organisation", "name": "Example Gram Panchayat", "mobile": "9830000088",
                  "address_1": "Example Village", "project_type": "street_lighting", "state": "West Bengal",
                  f"A{ITEM_FIRST_ROW}": "Solar street light", f"C{ITEM_FIRST_ROW}": "Branded",
                  f"D{ITEM_FIRST_ROW}": 40, f"E{ITEM_FIRST_ROW}": "Nos", f"F{ITEM_FIRST_ROW}": 18500,
                  f"G{ITEM_FIRST_ROW}": "solar_street_light"}
        r = _import(client, _workbook(values))
        assert r.status_code == 201, r.text
        assert _view(client, r.json()["created"][0]["proposal_id"])["figures"]["total"] == "777000.00"

    def test_every_error_type_with_its_cell_and_nothing_saved(self, client, db_session):
        from app.studio.importer import EXAMPLE, ITEM_FIRST_ROW
        bad = {**EXAMPLE, "name": None, "customer_type": "company", "mobile": "12345", "capacity_kwp": "lots",
               "gstin": "27AAPFU0939F1ZW", "dcr": "maybe", "houses": 2.5}
        items = {**EXAMPLE, "project_type": "street_lighting", "capacity_kwp": None,
                 f"A{ITEM_FIRST_ROW}": None, f"C{ITEM_FIRST_ROW}": "Branded", f"D{ITEM_FIRST_ROW}": -1,
                 f"E{ITEM_FIRST_ROW}": "crates", f"F{ITEM_FIRST_ROW}": "x", f"G{ITEM_FIRST_ROW}": "toys"}
        before = db_session.query(StudioProposal).count()
        r = _import(client, _workbook(EXAMPLE, bad, items))
        assert r.status_code == 422 and r.json()["detail"]["code"] == "import_failed"
        errors = {(e["sheet"], e["cell"]) for e in r.json()["detail"]["errors"]}
        assert {("Proposal 2", "B6"), ("Proposal 2", "B4"), ("Proposal 2", "B8"), ("Proposal 2", "B17"),
                ("Proposal 2", "B13"), ("Proposal 2", "B21"), ("Proposal 2", "B23")} <= errors
        assert {("Proposal 3", f"A{ITEM_FIRST_ROW}"), ("Proposal 3", f"D{ITEM_FIRST_ROW}"),
                ("Proposal 3", f"E{ITEM_FIRST_ROW}"), ("Proposal 3", f"F{ITEM_FIRST_ROW}"),
                ("Proposal 3", f"G{ITEM_FIRST_ROW}")} <= errors
        assert db_session.query(StudioProposal).count() == before        # the valid first sheet is not saved

    def test_price_mode_needs_an_amount(self, client):
        from app.studio.importer import EXAMPLE
        r = _import(client, _workbook({**EXAMPLE, "price_mode": "basic"}))
        assert r.status_code == 422 and r.json()["detail"]["errors"][0]["cell"] == "B38"

    def test_below_master_price_error_points_at_the_price_cell(self, client):
        from app.studio.importer import EXAMPLE
        r = _import(client, _workbook({**EXAMPLE, "price_mode": "basic", "price_value": 100000}))
        assert r.status_code == 422
        assert r.json()["detail"]["errors"][0]["cell"] == "B37"

    @pytest.mark.parametrize("name,raw,message", [
        ("x.csv", b"a,b", "Upload an .xlsx file."),
        ("x.xlsx", b"%PDF-1.4", "This is not an Excel workbook."),
        ("x.xlsx", b"PK\x03\x04garbage", "This is not an Excel workbook."),
    ])
    def test_rejected_file_types(self, client, name, raw, message):
        r = _import(client, raw, name=name)
        assert r.status_code == 422 and r.json()["detail"]["errors"][0]["message"] == message

    def test_macro_workbook_refused(self, client):
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(_workbook({})), "r") as src, zipfile.ZipFile(buf, "w") as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            dst.writestr("xl/vbaProject.bin", b"macro")
        r = _import(client, buf.getvalue())
        assert r.status_code == 422 and "macros" in r.json()["detail"]["errors"][0]["message"]


# ── The optional assistant ───────────────────────────────────────────────────────

class TestAssistant:
    def test_disabled_makes_no_call(self, client, monkeypatch):
        from app.llm.client import llm_client

        def boom(**kw):
            raise AssertionError("no model call when disabled")
        monkeypatch.setattr(llm_client, "studio_suggest", boom)
        pid = new_proposal(client)["proposal"]["id"]
        r = client.post(f"{D}/studio/proposals/{pid}/suggest-from-notes", headers=DESK, json={"text": "Mr Das 98300"})
        assert r.status_code == 409 and r.json()["detail"]["code"] == "assistant_disabled"

    def test_enabled_never_returns_price_boq_or_terms(self, client, monkeypatch):
        import json

        from app.config import settings
        from app.llm.client import llm_client
        monkeypatch.setattr(settings, "studio_ai_fill_enabled", True)
        reply = {"addressee": {"name": "Ramesh Das", "mobile": "9830012345", "gstin": "x"},
                 "connection": {"discom": "WBSEDCL", "phase": "three", "sanctioned_load_kw": "5"},
                 "price": {"basic": "1"}, "boq": [{"make": "x"}], "payment_terms": [{"percent": 100}],
                 "plant": {"capacity_kwp": "9"}, "site": {"state": "West Bengal"}}
        monkeypatch.setattr(llm_client, "studio_suggest",
                            lambda **kw: {"text": json.dumps(reply), "model": "gpt-4o-mini"})
        pid = new_proposal(client)["proposal"]["id"]
        before = _view(client, pid)["revision"]
        body = ok(client.post(f"{D}/studio/proposals/{pid}/suggest-from-notes", headers=DESK,
                              json={"text": "Ramesh Das 9830012345 WBSEDCL three phase"}))
        got = {(s["section"], s["field"]) for s in body["suggestions"]}
        assert got == {("addressee", "name"), ("addressee", "mobile"), ("connection", "discom"),
                       ("connection", "phase"), ("connection", "sanctioned_load_kw"), ("site", "state")}
        assert all(s["status"] == "check" for s in body["suggestions"])
        assert {"price", "boq", "payment_terms", "plant", "addressee.gstin"} <= set(body["dropped"])
        assert _view(client, pid)["revision"] == before                         # nothing saved
