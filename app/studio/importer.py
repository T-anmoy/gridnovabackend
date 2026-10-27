"""Fast input from an Excel sheet (BE-18 §11).

    template()     the workbook staff fill: one input sheet ("Proposal") with labelled cells
                   and drop-downs, a "Read me" sheet and a filled "Example" sheet.
    load(raw)      every filled input sheet (any sheet whose name starts with "Proposal") →
                   rows of cleaned values, or field-level errors with the cell reference.
    create(…)      one draft per sheet, all or nothing: if any sheet fails, nothing is saved.

Values are read, never evaluated (data_only workbook, formulas ignored). Only .xlsx up to
MAX_BYTES is accepted.
"""
from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from app.studio import config as studio_config

MAX_BYTES = 2 * 1024 * 1024
MAX_SHEETS = 20
INPUT_PREFIX = "Proposal"
ITEM_FIRST_ROW, ITEM_ROWS = 52, 20
YES_NO = ("yes", "no")
PHASES = ("single", "three")
PRICE_MODES = ("price_list", "basic", "rate_per_kwp", "total")


class ImportError_(ValueError):
    def __init__(self, errors: list[dict]):
        super().__init__("import failed")
        self.errors = errors


def _discoms() -> list[str]:
    from app.survey import config as survey_config
    return list(survey_config.load()["discom_options"])


# (row, key, label, kind, options/help). Column A label, B value, C help.
def fields() -> list[tuple]:
    from app.pricing import price_list
    from app.proposals import document as d
    from app.proposals import document_config
    return [
        (3, None, "CUSTOMER", "heading", None),
        (4, "customer_type", "Customer type", "choice", ["individual", "organisation"]),
        (5, "title", "Title", "choice_optional", document_config.load()["titles"]),
        (6, "name", "Customer / organisation name", "text", "Required"),
        (7, "attention", "Contact person (organisation)", "text_optional", "e.g. Mr. A. Kumar (Director)"),
        (8, "mobile", "Mobile (10 digits)", "mobile", "Required"),
        (9, "email", "Email", "text_optional", "For emailing the proposal"),
        (10, "address_1", "Billing address line 1", "text", "Required"),
        (11, "address_2", "Billing address line 2", "text_optional", None),
        (12, "address_3", "Billing address line 3", "text_optional", None),
        (13, "gstin", "GSTIN", "text_optional", "Optional; checked"),
        (15, None, "PROJECT", "heading", None),
        (16, "project_type", "Project type", "choice", list(studio_config.PROJECT_TYPES)),
        (17, "capacity_kwp", "Capacity (kWp)", "decimal_optional", "Not for itemised supply"),
        (18, "phase", "Electrical supply", "choice_optional", list(PHASES)),
        (19, "package", "Package (up to 10 kWp)", "choice_optional", list(price_list.PACKAGES)),
        (20, "module_technology", "Module technology", "choice_optional", d.technologies()),
        (21, "dcr", "DCR modules", "yes_no", None),
        (22, "subsidy_opted", "PM Surya Ghar subsidy opted", "yes_no", None),
        (23, "houses", "Number of houses (housing society)", "int_optional", None),
        (24, "roof_area_sqft", "Shadow-free roof area (sq. ft.)", "decimal_optional", None),
        (26, None, "SITE & CONNECTION", "heading", None),
        (27, "site_1", "Site address line 1", "text_optional", "Blank: the billing address"),
        (28, "site_2", "Site address line 2", "text_optional", None),
        (29, "town", "Town (named in the letter)", "text_optional", None),
        (30, "state", "State", "text_optional", "e.g. West Bengal"),
        (31, "discom", "DISCOM", "choice_optional", _discoms()),
        (32, "consumer_number", "Consumer number", "text_optional", None),
        (33, "sanctioned_load_kw", "Sanctioned load (kW)", "decimal_optional", None),
        (34, "category", "Consumer category", "choice_optional", list(d.CATEGORIES)),
        (36, None, "PRICE", "heading", None),
        (37, "price_mode", "Price from", "choice_optional", list(PRICE_MODES)),
        (38, "price_value", "Amount (Rs.) — basic, per kWp or total", "decimal_optional", "Not for price_list"),
        (40, None, "SAVINGS (from the customer's bill, optional)", "heading", None),
        (41, "monthly_units", "Average monthly units", "decimal_optional", None),
        (42, "monthly_bill", "Average monthly bill (Rs.)", "decimal_optional", None),
        (43, "tariff", "Tariff (Rs. per unit)", "decimal_optional", "Needed for the savings page"),
        (45, None, "NOTES", "heading", None),
        (46, "note_1", "Special note 1", "text_optional", None),
        (47, "note_2", "Special note 2", "text_optional", None),
        (48, "note_3", "Special note 3", "text_optional", None),
    ]


ITEM_COLUMNS = (("A", "description"), ("B", "specification"), ("C", "make"), ("D", "quantity"), ("E", "unit"),
                ("F", "unit_rate"), ("G", "category"))


def template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = ["Grid Nova Energy — Proposal Studio import sheet",
             "",
             "1. Fill the sheet called 'Proposal' — one sheet makes one draft proposal.",
             "2. For more proposals, copy the 'Proposal' sheet; keep each copy's name starting with 'Proposal'.",
             "3. Use the drop-downs where a cell has one. Leave a cell blank when you do not know the value.",
             "4. Itemised supply (street lighting): fill the item table from row 52 instead of the capacity.",
             "5. Upload the file in the Proposal Studio. Errors are listed with their cell; nothing is saved until "
             "every sheet is correct.",
             "6. The proposal is then checked and priced in the Studio as usual. Prices are always recomputed: "
             "never type GST or totals."]
    for i, line in enumerate(lines, 1):
        readme.cell(row=i, column=1, value=line)
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 110
    for title, example in (("Proposal", False), ("Example", True)):
        ws = wb.create_sheet(title)
        _input_sheet(ws, example, DataValidation, Font, PatternFill)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


EXAMPLE = {"customer_type": "individual", "title": "Mr.", "name": "Test Kumar Das", "mobile": "9830000001",
           "address_1": "12 Example Lane, Habra", "address_2": "North 24 Parganas, West Bengal – 743263",
           "project_type": "residential_rooftop", "capacity_kwp": 3, "phase": "single", "package": "NOVA_ASSURED",
           "module_technology": "Mono PERC", "dcr": "yes", "subsidy_opted": "yes", "town": "Habra",
           "state": "West Bengal", "discom": "WBSEDCL", "consumer_number": "100000001", "category": "Residential",
           "price_mode": "price_list"}


def _input_sheet(ws, example: bool, DataValidation, Font, PatternFill) -> None:
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 40
    ws["A1"] = "Grid Nova Energy — Proposal Studio input" + (" (EXAMPLE — not imported)" if example else "")
    ws["A1"].font = Font(bold=True, size=13)
    fill = PatternFill("solid", fgColor="FFF7E8")
    for row, key, label, kind, extra in fields():
        ws.cell(row=row, column=1, value=label).font = Font(bold=kind == "heading")
        if kind == "heading":
            continue
        cell = ws.cell(row=row, column=2)
        cell.fill = fill
        if example and key in EXAMPLE:
            cell.value = EXAMPLE[key]
        options = extra if isinstance(extra, list) else (list(YES_NO) if kind == "yes_no" else None)
        if options:
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(cell.coordinate)
        elif isinstance(extra, str):
            ws.cell(row=row, column=3, value=extra)
    ws.cell(row=ITEM_FIRST_ROW - 2, column=1, value="ITEMS (itemised supply only)").font = Font(bold=True)
    cats = list(studio_config.load()["itemised"]["categories"])
    units = list(studio_config.load()["itemised"]["units"])
    for col, key in ITEM_COLUMNS:
        ws[f"{col}{ITEM_FIRST_ROW - 1}"] = key.replace("_", " ").title()
        ws[f"{col}{ITEM_FIRST_ROW - 1}"].font = Font(bold=True)
    for options, col in ((units, "E"), (cats, "G")):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}{ITEM_FIRST_ROW}:{col}{ITEM_FIRST_ROW + ITEM_ROWS - 1}")


# ── Reading ─────────────────────────────────────────────────────────────────────

def _err(errors: list, sheet: str, cell: str, field: str, message: str) -> None:
    errors.append({"sheet": sheet, "cell": cell, "field": field, "message": message})


def _text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = " ".join(str(value).split())
    return text or None


def _decimal(value) -> Decimal:
    if isinstance(value, bool):
        raise InvalidOperation
    return Decimal(str(value).replace(",", "").strip())


def check_file(raw: bytes, filename: str, content_type: Optional[str]) -> None:
    if not filename.lower().endswith(".xlsx"):
        raise ImportError_([{"sheet": None, "cell": None, "field": "file", "message": "Upload an .xlsx file."}])
    if len(raw) > MAX_BYTES:
        raise ImportError_([{"sheet": None, "cell": None, "field": "file", "message": "The file is larger than 2 MB."}])
    if not raw.startswith(b"PK"):
        raise ImportError_([{"sheet": None, "cell": None, "field": "file", "message": "This is not an Excel workbook."}])
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if "xl/workbook.xml" not in z.namelist():
                raise zipfile.BadZipFile
            if sum(i.file_size for i in z.infolist()) > 40 * MAX_BYTES:
                raise ImportError_([{"sheet": None, "cell": None, "field": "file",
                                     "message": "The workbook expands too large."}])
            if any(n.startswith("xl/vbaProject") for n in z.namelist()):
                raise ImportError_([{"sheet": None, "cell": None, "field": "file",
                                     "message": "Workbooks with macros are not accepted."}])
    except zipfile.BadZipFile:
        raise ImportError_([{"sheet": None, "cell": None, "field": "file", "message": "This is not an Excel workbook."}])


def load(raw: bytes) -> list[dict]:
    """[{sheet, values, items}] for every filled input sheet. Raises ImportError_ with every
    error found (cell references included)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = [ws for ws in wb.worksheets if ws.title.startswith(INPUT_PREFIX)][:MAX_SHEETS + 1]
    errors: list[dict] = []
    if len(sheets) > MAX_SHEETS:
        raise ImportError_([{"sheet": None, "cell": None, "field": "file",
                             "message": f"At most {MAX_SHEETS} proposal sheets per file."}])
    out = []
    for ws in sheets:
        grid = {}
        for row in ws.iter_rows(min_row=1, max_row=ITEM_FIRST_ROW + ITEM_ROWS, max_col=7):
            for cell in row:
                if cell.value is not None and hasattr(cell, "coordinate"):
                    grid[cell.coordinate] = cell.value
        if not any(grid.get(f"B{row}") is not None for row, key, *_ in fields() if key):
            continue                                           # an empty copy: skipped
        values, sheet_errors = _read_fields(ws.title, grid)
        items = _read_items(ws.title, grid, sheet_errors)
        _cross_checks(ws.title, values, items, sheet_errors)
        errors += sheet_errors
        out.append({"sheet": ws.title, "values": values, "items": items})
    wb.close()
    if not out and not errors:
        errors.append({"sheet": None, "cell": None, "field": "file", "message": "No filled 'Proposal' sheet found."})
    if errors:
        raise ImportError_(errors)
    return out


def _read_fields(sheet: str, grid: dict) -> tuple[dict, list]:
    values: dict[str, Any] = {}
    errors: list = []
    for row, key, label, kind, extra in fields():
        if key is None:
            continue
        cell = f"B{row}"
        raw = grid.get(cell)
        text = _text(raw)
        optional = kind.endswith("_optional") or kind == "yes_no"
        if text is None:
            if not optional:
                _err(errors, sheet, cell, key, f"{label} is required.")
            values[key] = False if kind == "yes_no" else None
            continue
        base = kind.replace("_optional", "")
        if base == "choice":
            if text not in extra:
                _err(errors, sheet, cell, key, f"{label}: choose one of {', '.join(extra)}.")
            values[key] = text
        elif base == "yes_no":
            if text.lower() not in YES_NO:
                _err(errors, sheet, cell, key, f"{label}: yes or no.")
            values[key] = text.lower() == "yes"
        elif base == "mobile":
            digits = re.sub(r"\D", "", text)
            digits = digits[2:] if len(digits) == 12 and digits.startswith("91") else digits
            if not re.fullmatch(r"[6-9]\d{9}", digits):
                _err(errors, sheet, cell, key, f"{label}: a 10-digit Indian mobile number.")
            values[key] = digits
        elif base in ("decimal", "int"):
            try:
                number = _decimal(raw)
                if number <= 0 or (base == "int" and number != number.to_integral_value()):
                    raise InvalidOperation
                values[key] = int(number) if base == "int" else number
            except (InvalidOperation, ValueError):
                _err(errors, sheet, cell, key, f"{label}: a positive {'whole ' if base == 'int' else ''}number.")
                values[key] = None
        else:
            if len(text) > 200:
                _err(errors, sheet, cell, key, f"{label}: at most 200 characters.")
            values[key] = text
    if values.get("gstin"):
        from app.studio.checks import gstin_valid
        if not gstin_valid(values["gstin"].upper()):
            _err(errors, sheet, f"B{_row('gstin')}", "gstin", "GSTIN: not valid (format or check digit).")
    return values, errors


def _row(key: str) -> int:
    return next(row for row, k, *_ in fields() if k == key)


def _read_items(sheet: str, grid: dict, errors: list) -> list[dict]:
    cats = studio_config.load()["itemised"]["categories"]
    units = studio_config.load()["itemised"]["units"]
    items = []
    for r in range(ITEM_FIRST_ROW, ITEM_FIRST_ROW + ITEM_ROWS):
        row = {key: grid.get(f"{col}{r}") for col, key in ITEM_COLUMNS}
        if all(v is None for v in row.values()):
            continue
        item = {"description": _text(row["description"]), "specification": _text(row["specification"]),
                "make": _text(row["make"]) or "–", "unit": _text(row["unit"]), "category": _text(row["category"])}
        if not item["description"]:
            _err(errors, sheet, f"A{r}", "items.description", "Item description is required.")
        for col, key in (("D", "quantity"), ("F", "unit_rate")):
            try:
                number = _decimal(row[key])
                if number <= 0:
                    raise InvalidOperation
                item[key] = str(number)
            except (InvalidOperation, ValueError, TypeError):
                _err(errors, sheet, f"{col}{r}", f"items.{key}", f"{key.replace('_', ' ').title()}: a positive number.")
        if item["unit"] not in units:
            _err(errors, sheet, f"E{r}", "items.unit", f"Unit: one of {', '.join(units)}.")
        if item["category"] not in cats:
            _err(errors, sheet, f"G{r}", "items.category", f"Category: one of {', '.join(cats)}.")
        items.append({**item, "_row": r})
    return items


def _cross_checks(sheet: str, values: dict, items: list, errors: list) -> None:
    ptype = values.get("project_type")
    itemised = ptype in studio_config.PROJECT_TYPES and studio_config.project_type(ptype)["layout"] == "itemised"
    if itemised and not items:
        _err(errors, sheet, f"A{ITEM_FIRST_ROW}", "items", "An itemised proposal needs at least one item.")
    if not itemised and ptype and not values.get("capacity_kwp"):
        _err(errors, sheet, f"B{_row('capacity_kwp')}", "capacity_kwp", "Capacity (kWp) is required.")
    mode = values.get("price_mode")
    if mode and mode != "price_list" and not values.get("price_value"):
        _err(errors, sheet, f"B{_row('price_value')}", "price_value", "Give the amount for this price mode.")


# ── Creating the drafts ─────────────────────────────────────────────────────────

SECTION_CELL = {"addressee": "name", "site": "site_1", "connection": "discom", "plant": "capacity_kwp",
                "system": "module_technology", "pm_surya_ghar": "subsidy_opted", "project": "houses",
                "savings": "tariff", "notes_free": "note_1", "items": None, "price": "price_mode", "display": "tariff"}


def create(db, rows: list[dict], *, actor: str, can_price: bool) -> list[dict]:
    """One draft per sheet, inside the caller's transaction; the caller rolls back on error."""
    from app.studio import service
    errors, created = [], []
    for row in rows:
        sheet, v, items = row["sheet"], row["values"], row["items"]
        if v.get("price_mode") and not can_price:
            _err(errors, sheet, f"B{_row('price_mode')}", "price_mode", "Your role cannot set prices; leave it blank.")
            continue
        try:
            with db.begin_nested():
                created.append(_one(db, service, v, items, actor))
        except service.StudioError as exc:
            section = (exc.field or "").split(".")[0].split("[")[0]
            key = SECTION_CELL.get(section)
            cell = (f"A{items[0]['_row']}" if section == "items" and items else
                    f"B{_row(key)}" if key else None)
            _err(errors, sheet, cell, exc.field or section, exc.message)
    if errors:
        raise ImportError_(errors)
    return created


def _one(db, service, v: dict, items: list[dict], actor: str) -> dict:
    customer, dups = service.create_customer(db, {
        "customer_type": v["customer_type"], "title": v.get("title"), "name": v["name"], "attention": v.get("attention"),
        "mobile": v["mobile"], "email": v.get("email"),
        "address_lines": [x for x in (v.get("address_1"), v.get("address_2"), v.get("address_3")) if x],
        "gstin": (v.get("gstin") or "").upper() or None}, actor)
    p = service.create_proposal(db, customer=customer, project_type=v["project_type"], owner_id=actor, actor=actor)

    def put(section, value):
        doc = service.doc_of(service.version_of(db, p))
        service.edit(db, p, section=section, value=value, base_revision=doc.revision, actor=actor, can_price=True)

    itemised = studio_config.project_type(v["project_type"])["layout"] == "itemised"
    site_lines = [x for x in (v.get("site_1"), v.get("site_2")) if x] or list(customer.address_lines)
    put("site", {"address_lines": site_lines, "town": v.get("town"), "state": v.get("state")})
    if not itemised:
        put("connection", {"discom": v.get("discom"), "consumer_number": v.get("consumer_number"),
                           "sanctioned_load_kw": str(v["sanctioned_load_kw"]) if v.get("sanctioned_load_kw") else None,
                           "category": v.get("category") or studio_config.project_type(v["project_type"])["category"],
                           "phase": v.get("phase")})
        if v.get("module_technology"):
            put("system", {"module_technology": v["module_technology"]})
        put("plant", {"capacity_kwp": str(v["capacity_kwp"]), "package": v.get("package"), "dcr": v.get("dcr")})
        if v.get("subsidy_opted"):
            put("pm_surya_ghar", {"subsidy_opted": True})
        if v.get("houses") or v.get("roof_area_sqft"):
            put("project", {"houses": v.get("houses"),
                            "roof_area_sqft": str(v["roof_area_sqft"]) if v.get("roof_area_sqft") else None})
        if v.get("tariff"):
            put("savings", {"monthly_units": str(v["monthly_units"]) if v.get("monthly_units") else None,
                            "monthly_bill": str(v["monthly_bill"]) if v.get("monthly_bill") else None,
                            "tariff_per_unit": str(v["tariff"])})
            put("display", {"show_savings": True})
    else:
        put("items", [{k: x[k] for k in ("description", "specification", "make", "quantity", "unit", "unit_rate",
                                         "category")} for x in items])
    notes = [x for x in (v.get("note_1"), v.get("note_2"), v.get("note_3")) if x]
    if notes:
        put("notes_free", notes)
    mode = v.get("price_mode")
    if mode and not itemised:
        value = {"price_list": {"mode": "price_list"}, "basic": {"mode": "staff_basic", "basic": str(v["price_value"])},
                 "rate_per_kwp": {"mode": "staff_basic", "rate_per_kwp": str(v["price_value"])},
                 "total": {"mode": "staff_total", "total": str(v["price_value"])}}[mode]
        put("price", value)
    service.event(db, p, "imported", version=1, actor=actor, detail={"duplicates": dups})
    return {"proposal_id": p.id, "code": p.code, "title": p.title, "duplicate_customers": dups}
